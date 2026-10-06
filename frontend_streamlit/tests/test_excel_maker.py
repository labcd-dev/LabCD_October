import json
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from frontend_streamlit import excel_maker as maker
from frontend_streamlit import conversation_core as core


@pytest.mark.parametrize("text,expected_headers,expected_rows", [
    ("Make Excel:\nName,Score,Postal Code\nAli,12.5,00120\nMina,9,00440",
     ["Name", "Score", "Postal Code"], [["Ali", 12.5, "00120"], ["Mina", 9, "00440"]]),
    ("| State | Derivative |\n| --- | ---: |\n| s_angle | 1.25 |\n| s_speed | -0.5 |",
     ["State", "Derivative"], [["s_angle", 1.25], ["s_speed", -0.5]]),
    ('[{"sensor":"a","value":2},{"sensor":"b","value":3.5}]',
     ["sensor", "value"], [["a", 2], ["b", 3.5]]),
])
def test_parse_raw_data_formats(text, expected_headers, expected_rows):
    assert maker.parse_raw_table(text) == (expected_headers, expected_rows)


def test_parse_flattened_step_response_table_from_chat_text():
    values = []
    for index in range(41):
        values.extend((str(index), f"{index * 0.05:.2f}", "1.000" if index < 30 else "0.000",
                       f"{[0.0000, 0.0029, 0.0113, 0.0243, 0.0415][index] if index < 5 else 0.05 + index / 50:.4f}"))
    text = "k t (s) u(k) Input y(k) Output " + " ".join(values)

    headers, rows = maker.parse_raw_table(text)

    assert headers == ["k", "t (s)", "u(k) Input", "y(k) Output"]
    assert len(rows) == 41
    assert rows[:2] == [[0, 0.0, 1, 0.0], [1, 0.05, 1, 0.0029]]
    assert rows[-1][:3] == [40, 2.0, 0]


@pytest.mark.parametrize("question", [
    "Make me an Excel file from this table",
    "Create a spreadsheet please",
    "Export this as xlsx",
    "Excel Maker, use the data above",
])
def test_excel_creation_requests_are_recognized(question):
    assert maker.is_excel_request(question)


def test_general_excel_question_does_not_start_workbook_creation():
    assert not maker.is_excel_request("What is the difference between Excel and CSV?")


def test_create_workbook_preserves_data_and_stores_formula_like_text_safely(tmp_path):
    raw = ('Make an Excel workbook:\nName,Postal Code,Score,Note\n'
           'Ali,00120,0.25,"=HYPERLINK(""https://example.test"",""open"")"')
    chat = core.new_chat()
    chat["messages"].append({"role": "user", "content": raw})

    class Client:
        settings = SimpleNamespace(provider="openai", model="configured-current-model")
        calls = []

        def complete(self, system, user):
            self.calls.append((system, user))
            request = json.loads(user)
            assert request["columns"] == ["Name", "Postal Code", "Score", "Note"]
            assert request["data_row_count"] == 1
            assert "HYPERLINK" not in user
            assert "Ali" not in user
            assert "Excel Maker" in system
            return json.dumps({"title": "Sensor export", "sheet_name": "Raw signals"})

    client = Client()
    result = maker.create_workbook(chat, "Make an Excel workbook from my raw data", client=client,
                                   output_dir=tmp_path)

    path = result["generated_artifact"]["path"]
    assert result["generated_artifact"]["rows"] == 1
    assert result["generated_artifact"]["columns"] == 4
    assert result["model"] == "configured-current-model"
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Raw signals"]
    assert sheet.freeze_panes == "A5"
    assert sheet["A1"].value == "Sensor export"
    assert sheet["A4"].value == "Name"
    assert sheet["B5"].value == "00120"
    assert sheet["B5"].data_type == "s"
    assert sheet["C5"].value == 0.25
    assert sheet["D5"].value == '=HYPERLINK("https://example.test","open")'
    assert sheet["D5"].data_type == "s"
    assert list(sheet.tables) == ["ChatData"]
    assert len(client.calls) == 1


def test_excel_maker_can_use_raw_table_from_earlier_chat_message(tmp_path):
    chat = core.new_chat()
    chat["messages"] = [
        {"role": "user", "content": "time,position\n0,1.2\n0.1,1.5"},
        {"role": "assistant", "content": "What would you like to do with these measurements?"},
        {"role": "user", "content": "Make an Excel workbook from that."},
    ]

    class Client:
        def complete(self, system, user):
            return '{"title":"Position data","sheet_name":"Measurements"}'

    result = maker.create_workbook(chat, "Make an Excel workbook from that", client=Client(), output_dir=tmp_path)
    sheet = load_workbook(result["generated_artifact"]["path"], data_only=True)["Measurements"]
    assert sheet["A4"].value == "time"
    assert sheet["A5"].value == 0
    assert sheet["B6"].value == 1.5


def test_excel_maker_explains_when_chat_has_no_tabular_data():
    with pytest.raises(ValueError, match="Paste a header row"):
        maker.create_workbook(core.new_chat(), "Make an Excel workbook", client=object())


def test_malformed_table_rows_are_not_silently_dropped():
    with pytest.raises(ValueError, match="different numbers of cells"):
        maker.parse_raw_table("Name,Score\nAli,1\nMina,2,unexpected")


def test_timezone_offsets_and_unrepresentable_numbers_stay_as_text():
    timestamp = "2026-09-30T12:30:00+03:30"
    assert maker._typed_value(timestamp, "timestamp") == timestamp
    assert maker._typed_value("1e999", "measurement") == "1e999"


def test_model_supplied_title_cannot_become_an_excel_formula(tmp_path):
    path = tmp_path / "safe-title.xlsx"
    maker.build_workbook(path, "=HYPERLINK(\"https://example.test\")", "Data", ["Value"], [[1]])
    title = load_workbook(path, data_only=False).active["A1"]
    assert title.data_type == "s"
    assert title.value.startswith("=HYPERLINK")


def test_latest_uploaded_file_takes_precedence_over_older_pasted_table(tmp_path):
    uploaded = tmp_path / "new_measurements.csv"
    uploaded.write_text("state,error\nspeed,0.02\nangle,0.04\n", encoding="utf-8")
    chat = core.new_chat()
    chat["messages"] = [
        {"role": "user", "content": "time,value\n0,1\n1,2"},
        {"role": "user", "content": "Make an Excel workbook from this upload.",
         "attachment": {"path": str(uploaded)}},
    ]
    headers, rows = maker._source_table(chat)
    assert headers == ["state", "error"]
    assert rows == [["speed", 0.02], ["angle", 0.04]]


def test_unreadable_latest_attachment_does_not_fall_back_to_stale_table(tmp_path):
    uploaded = tmp_path / "empty.csv"
    uploaded.write_text("", encoding="utf-8")
    chat = core.new_chat()
    chat["messages"] = [
        {"role": "user", "content": "time,value\n0,1\n1,2"},
        {"role": "user", "content": "Make an Excel workbook from this upload.",
         "attachment": {"path": str(uploaded)}},
    ]
    with pytest.raises(ValueError, match="latest attachment"):
        maker._source_table(chat)
